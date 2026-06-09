#!/usr/bin/env python3
"""Generate presentation slides for the portfolio project."""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

OUTPUT = Path(__file__).resolve().parent.parent / "presentation.pptx"


def add_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_content_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    tf = slide.placeholders[1].text_frame
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(18)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "Supply Chain Intelligence\n& Demand Forecasting Platform",
        "Portfolio Project | Data Analytics & Engineering",
    )

    add_content_slide(prs, "Business Problem", [
        "Retail company with multiple warehouses, 100+ suppliers, 10K+ products",
        "Need centralized analytics for inventory, suppliers, and logistics",
        "Goals: Prevent stock-outs, optimize inventory, forecast demand",
        "Target users: Supply chain managers, analysts, executives",
    ])

    add_content_slide(prs, "Solution Overview", [
        "End-to-end analytics platform with ETL, KPIs, ML forecasting",
        "520K+ order transactions across 25 warehouses",
        "Automated insights and inventory optimization engine",
        "Executive dashboards and REST API for integration",
    ])

    add_content_slide(prs, "Tech Stack", [
        "Backend: Python, FastAPI | Database: PostgreSQL",
        "Data Processing: Pandas, NumPy",
        "ML: Scikit-learn, XGBoost, Facebook Prophet",
        "BI: Power BI | Deployment: Docker, Docker Compose",
    ])

    add_content_slide(prs, "Architecture", [
        "Phase 1: Synthetic data generation → ETL → PostgreSQL",
        "Phase 2: KPI Engine (Inventory, Supplier, Logistics, Service)",
        "Phase 3: 40+ SQL analytics queries with window functions",
        "Phase 4-6: Inventory, Supplier, Logistics intelligence modules",
        "Phase 7-8: Demand forecasting + Inventory optimization (EOQ)",
        "Phase 9-10: Power BI dashboards + Automated insights engine",
    ])

    add_content_slide(prs, "Key Metrics & KPIs", [
        "Inventory: Turnover ratio, Days of inventory, Carrying cost",
        "Supplier: On-time delivery %, Lead time, Reliability score",
        "Logistics: Delay rate, Transportation cost, Route efficiency",
        "Service: Fill rate, Stock-out frequency, Order accuracy",
    ])

    add_content_slide(prs, "Demand Forecasting", [
        "4 Models: Linear Regression, Random Forest, XGBoost, Prophet",
        "Horizons: 3, 6, and 12 months",
        "Evaluation: RMSE, MAE, MAPE",
        "Best model selected automatically for production forecasts",
    ])

    add_content_slide(prs, "Results & Impact", [
        "500K+ transactions processed with full data quality pipeline",
        "Automated reorder recommendations with EOQ and safety stock",
        "Supplier scorecards with A-D grading and rankings",
        "Actionable insights for stock-out prevention and route optimization",
    ])

    add_content_slide(prs, "Skills Demonstrated", [
        "SQL (Advanced): CTEs, Window Functions, Cohort Analysis",
        "Python: ETL, Analytics, ML, API Development",
        "Data Engineering: Pipeline design, Schema optimization",
        "Business Intelligence: KPI design, Dashboard specifications",
        "DevOps: Docker containerization, CI-ready structure",
    ])

    add_title_slide(prs, "Thank You", "GitHub: srimaan07/Supply-Chain-Intelligence-Demand-Forecasting-Platform")

    prs.save(OUTPUT)
    print(f"Presentation saved to {OUTPUT}")


if __name__ == "__main__":
    main()
